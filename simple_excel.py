"""
Simple Excel operations for Robot Framework without RPA dependencies
"""
import openpyxl


def create_workbook(filename):
    """Create a new Excel workbook"""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    return workbook, worksheet, filename


def append_rows_to_worksheet(workbook, rows, header=False):
    """Append rows to the worksheet"""
    worksheet = workbook.active
    for row in rows:
        worksheet.append(row)
    return workbook


def save_workbook(workbook, filename):
    """Save the workbook"""
    workbook.save(filename)


def list_worksheets(workbook):
    """List all worksheets in the workbook"""
    return workbook.sheetnames


def read_worksheet(workbook, sheet_name=None, header=True):
    """Read worksheet data"""
    if sheet_name:
        ws = workbook[sheet_name]
    else:
        ws = workbook.active
        
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(cell is not None for cell in row):  # Skip empty rows
            rows.append(list(row))
            
    if header and rows:
        headers = rows[0]
        data_rows = rows[1:]
        result = []
        for row in data_rows:
            row_dict = {}
            for i, header in enumerate(headers):
                row_dict[header] = row[i] if i < len(row) else None
            result.append(row_dict)
        return result
    return rows


def set_cell_value(workbook, sheet_name, cell, value):
    """Set cell value"""
    if sheet_name:
        ws = workbook[sheet_name]
    else:
        ws = workbook.active
    ws[cell] = value


def close_workbook(workbook):
    """Close the workbook"""
    if workbook:
        workbook.close()


# Robot Framework compatible class
class SimpleExcel:
    ROBOT_LIBRARY_SCOPE = 'GLOBAL'
    
    def __init__(self):
        self.workbook = None
        self.filename = None
        
    @keyword("Append Rows To Worksheet")
    def append_rows_to_worksheet(self, rows, header=True):
        """Append rows to the worksheet"""
        if not self.worksheet:
            raise ValueError("No active worksheet. Create workbook first.")
            
        for row in rows:
            self.worksheet.append(row)
    
    @keyword("Save Workbook")        
    def save_workbook(self):
        """Save the workbook"""
        if not self.workbook:
            raise ValueError("No workbook to save.")
        self.workbook.save(self.filename)
        
    @keyword("List Worksheets")
    def list_worksheets(self):
        """List all worksheets in the workbook"""
        if not self.workbook:
            return []
        return self.workbook.sheetnames
        
    @keyword("Read Worksheet")
    def read_worksheet(self, sheet_name=None, header=True):
        """Read worksheet data"""
        if not self.workbook:
            raise ValueError("No workbook loaded.")
            
        if sheet_name:
            ws = self.workbook[sheet_name]
        else:
            ws = self.workbook.active
            
        rows = []
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):  # Skip empty rows
                rows.append(list(row))
                
        if header and rows:
            headers = rows[0]
            data_rows = rows[1:]
            result = []
            for row in data_rows:
                row_dict = {}
                for i, header in enumerate(headers):
                    row_dict[header] = row[i] if i < len(row) else None
                result.append(row_dict)
            return result
        return rows
        
    @keyword("Set Cell Value")
    def set_cell_value(self, sheet_name, cell, value):
        """Set cell value"""
        if sheet_name:
            ws = self.workbook[sheet_name]
        else:
            ws = self.workbook.active
        ws[cell] = value
        
    @keyword("Close Workbook")
    def close_workbook(self):
        """Close the workbook"""
        if self.workbook:
            self.workbook.close()
            self.workbook = None
            self.worksheet = None